"""Build the Problem 1 chapter from the 2026 official Word template and frozen CSVs.

This script writes paper artifacts only; it never imports or changes solver code.
"""

from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


PAPER = Path(__file__).resolve().parent.parent
ROOT = PAPER.parent.parent
FINAL = ROOT / "artifacts" / "problem1_inherited_fallback_c4140"
ORIGINAL = ROOT / "artifacts" / "problem1_full_c4140"
FIGURES = ROOT / "artifacts" / "problem1_paper_figures" / "figures"
WIDE = ROOT / "artifacts" / "problem1_paper_figures" / "data" / "appendix_problem1_wide.csv"
TEMPLATE = PAPER / "templates" / "官方模板转换.docx"
OUTPUT = PAPER / "draft" / "问题一章节草稿.docx"
MATH_FRAGMENTS = json.loads(
    (PAPER / "scripts" / "omml_fragments.json").read_text(encoding="utf-8")
)


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def east_asian_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")


def add_paragraph(doc, text, *, indent=True, align=None, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    east_asian_font(p.add_run(text), size=size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.0
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        east_asian_font(p.add_run(text), "黑体", 14, True)
    else:
        east_asian_font(p.add_run(text), "宋体", 12, True)
    return p


def add_equation(doc, latex, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    for child in etree.fromstring(MATH_FRAGMENTS[latex].encode("utf-8")):
        math.append(child)
    math_para.append(math)
    p._element.append(math_para)
    # Number belongs to the paragraph, while the equation itself stays editable OMML.
    east_asian_font(p.add_run(f"    （{number}）"), size=10)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = False
    east_asian_font(p.add_run(text), size=12)


def add_figure(doc, filename, caption, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURES / filename), width=Cm(width_cm))
    add_caption(doc, caption)


def _three_line_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single" if side in ("top", "bottom") else "nil")
        border.set(qn("w:sz"), "12" if side in ("top", "bottom") else "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if row_index == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                tc_borders = tc_pr.find(qn("w:tcBorders"))
                if tc_borders is None:
                    tc_borders = OxmlElement("w:tcBorders")
                    tc_pr.append(tc_borders)
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "4")
                bottom.set(qn("w:color"), "000000")
                tc_borders.append(bottom)
    return table


def _normalize_table_properties(table):
    """Keep table properties in the sequence required by WordprocessingML."""
    sequence = (
        "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
        "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
        "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
        "tblCaption", "tblDescription", "tblPrChange",
    )
    order = {qn(f"w:{name}"): index for index, name in enumerate(sequence)}
    properties = table._tbl.tblPr
    children = sorted(list(properties), key=lambda child: order.get(child.tag, len(order)))
    for child in children:
        properties.remove(child)
        properties.append(child)


def add_table(doc, headers, rows, caption, *, font_size=12, widths_cm=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    east_asian_font(p.add_run(caption), size=12)

    table = _three_line_table(doc, [headers, *rows])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths_cm is None
    if widths_cm is not None:
        for column, width in zip(table.columns, widths_cm):
            column.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths_cm):
                cell.width = Cm(width)
    _normalize_table_properties(table)
    # Repeat the header on continuation pages of the two 100-case appendix tables.
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_props.append(repeat)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = 1
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    east_asian_font(run, size=font_size, bold=(ri == 0))
        if ri == 0:
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next = True
    return table


def set_math_cell(cell, latex):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = OxmlElement("m:oMath")
    for child in etree.fromstring(MATH_FRAGMENTS[latex].encode("utf-8")):
        math.append(child)
    paragraph._element.append(math)


def main():
    final_rows = read_rows(FINAL / "problem1_inherited_fallback_results.csv")
    original_rows = read_rows(ORIGINAL / "problem1_full_results.csv")
    aggregate_rows = read_rows(FINAL / "problem1_inherited_fallback_aggregate.csv")
    appendix_rows = read_rows(WIDE)
    assert len(final_rows) == 500 and len(original_rows) == 400 and len(appendix_rows) == 100
    assert all(row["officially_measured"] == "True" for row in final_rows)
    assert all(row["status"] == "success" for row in original_rows)
    assert {int(row["cores"]) for row in final_rows} == set(range(1, 6))
    assert all(len([r for r in final_rows if r["case"] == f"case_{i:03d}"]) == 5 for i in range(1, 101))

    by_core = defaultdict(list)
    for row in final_rows:
        by_core[int(row["cores"])].append(row)
    aggregate = {int(row["cores"]): row for row in aggregate_rows}
    for k in range(1, 6):
        avg = sum(float(r["final_speedup"]) for r in by_core[k]) / 100
        assert math.isclose(avg, float(aggregate[k]["reported_arithmetic_mean_speedup"]), abs_tol=1e-12)
    final_lookup = {(r["case"], int(r["cores"])): r for r in final_rows}
    for row in appendix_rows:
        for k in range(1, 6):
            source = final_lookup[(row["case"], k)]
            assert int(row[f"makespan_k{k}"]) == int(source["final_makespan"])
            assert int(row[f"added_copy_bytes_k{k}"]) == int(source["final_added_copy_bytes"])
    improved = [r for r in final_rows if int(r["makespan_improvement"]) > 0]
    assert len(improved) == 36
    assert Counter(int(r["cores"]) for r in improved) == {3: 3, 4: 9, 5: 24}
    assert sum(int(r["added_copy_bytes_change"]) < 0 for r in improved) == 33
    assert sum(int(r["added_copy_bytes_change"]) == 0 for r in improved) == 3
    assert sum(r["winner"].startswith("b2a") for r in original_rows) == 245
    assert sum(r["final_winner"].startswith("b2a") for r in final_rows) == 220
    executed_times = sorted(float(r["actual_wall_time_seconds"]) for r in original_rows if r["run_disposition"] == "executed")
    assert len(executed_times) == 352
    assert math.isclose((executed_times[175] + executed_times[176]) / 2, 13.572004239, abs_tol=1e-9)
    example = final_lookup[("case_044", 5)]
    assert tuple(int(example[key]) for key in (
        "old_makespan", "final_makespan", "old_added_copy_bytes", "final_added_copy_bytes"
    )) == (154407, 87589, 4438112, 2915104)

    doc = Document(TEMPLATE)
    body_element = doc._element.body
    for child in list(body_element):
        if child.tag != qn("w:sectPr"):
            body_element.remove(child)
    # The official .doc references a logo through an absolute path on its
    # creator's computer. The interim chapter has no cover, so remove the
    # now-unused external image relation after clearing the template body.
    for relation_id, relation in list(doc.part.rels.items()):
        if relation.is_external and relation.reltype.endswith("/image"):
            del doc.part.rels[relation_id]
    # A standalone chapter starts at page 1; the submission template begins
    # its numbering after the cover and therefore carries start=0/titlePg.
    sect = doc.sections[0]._sectPr
    page_number = sect.find(qn("w:pgNumType"))
    if page_number is None:
        page_number = OxmlElement("w:pgNumType")
        sect.append(page_number)
    page_number.set(qn("w:start"), "1")
    title_page = sect.find(qn("w:titlePg"))
    if title_page is not None:
        sect.remove(title_page)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    add_heading(doc, "4 问题一：场景 A 下的多核切图与调度", 1)
    add_paragraph(doc, "场景 A 中，一个子图对应一个 Task；跨子图的数据由 DDR 周转。给定算子—张量有向无环图和核心数，问题一需要确定算子到子图的归属、子图到核心的分配，以及每个核心的子图执行顺序。本文以官方评测所得总完成时间为首要目标，并记录总附加数据搬运量，以反映切分边界和调度顺序带来的通信代价。")

    add_heading(doc, "4.1 图收缩、评价指标与符号", 2)
    add_paragraph(doc, "先按题目给定的合法性规则校验输入图，排除原始 COPY_IN、COPY_OUT 搬运算子，并收缩穿过这些算子的依赖路径，得到可划分计算算子的有向无环图。张量仍保留其标识、字节数及生产者—消费者关系，用于计算跨子图通信。每个可划分算子必须且只能落在一个子图内；子图之间的依赖与每个核心上的执行序列必须保持拓扑合法。")
    symbol_table = add_table(doc, ["符号", "含义", "单位"], [
        ["G=(V,E)", "收缩后的计算算子依赖图", "—"],
        ["k", "可用 NPU 核心数", "核"],
        ["ℓ(v)", "算子 v 的依赖层号", "层"],
        ["q_l", "第 l 层用于窗口切分的双管线权重", "cycle"],
        ["W_j", "由相邻完整依赖层形成的第 j 个窗口", "—"],
        ["S", "一个合法子图", "—"],
        ["p(S)", "子图 S 的双流水线代理工作量", "cycle"],
        ["b(S,T)", "商图边 S→T 的去重通信数据量", "byte"],
        ["r(S)", "子图 S 的通信加权上行秩", "cycle"],
        ["T_{i,k}", "案例 i、k 核方案的官方 Makespan", "cycle"],
        ["D_{i,k}", "案例 i、k 核方案的官方附加搬运量", "byte"],
        ["s_{i,k}", "案例 i、k 核相对于单核的加速比", "无量纲"],
    ], "表1 问题一主要符号", font_size=12, widths_cm=[3.0, 10.5, 3.0])
    for row, latex in zip(symbol_table.rows[1:], [
        r"G=(V,E)", r"k", r"\ell(v)", r"q_l", r"W_j", r"S", r"p(S)",
        r"b(S,T)", r"r(S)", r"T_{i,k}", r"D_{i,k}", r"s_{i,k}",
    ]):
        set_math_cell(row.cells[0], latex)
    add_paragraph(doc, "单核基准由官方评测得到。逐例加速比定义为同一案例单核完成时间与相应多核完成时间之比，再对 100 个案例取算术平均；这与先平均时间再取比值不同。总附加搬运量直接使用官方评测给出的字节数。")
    add_equation(doc, r"\bar{s}_k=\frac{1}{100}\sum_{i=1}^{100}\frac{T_{i,1}}{T_{i,k}}", "1")

    add_heading(doc, "4.2 对照切图与共用调度器", 2)
    add_paragraph(doc, "为了辨别结构切图的作用，设置两个确定性对照。B0 对合法拓扑序按算子工作量切成约 4k 个连续片段，代表只考虑顺序和粗粒度均衡的基本方案。B1 在 Kahn 拓扑构造中形成约 4k 个块：每块先选逆向关键路径优先级最高的就绪算子，随后优先选择与当前块共享张量字节最多的就绪算子，再以关键路径优先级和算子编号打破并列。B0、B1 均保持图的依赖合法性。")
    add_paragraph(doc, "三种切图统一进入同一商图调度器。对每个子图，以 Cube 对应的 M 管线周期和 Vector 对应的 V 管线周期之较大值作为计算时间代理；商图边的通信量对同一边上的张量 ID 去重，再计入直接算子依赖携带的字节。调度器据此计算通信加权的逆向上行秩，从就绪子图中选秩最高者，为其尝试每个核心并选择预计最早完成的核心。该“上行秩—最早完成时间”框架借鉴 HEFT [1]，本文仅将其任务粒度和成本模型改造到场景 A 的子图商图；它不是本文独立提出的调度原理。")
    add_equation(doc, r"p(S)=\max(\sum_{v\in S}c_M(v),\sum_{v\in S}c_V(v))", "2")
    add_equation(doc, r"r(S)=p(S)+\max_{T\in succ(S)}(\frac{b(S,T)}{60}+r(T))", "3")
    add_paragraph(doc, "式（3）中的 60 byte/cycle 是候选构造阶段的通信时间代理。对候选核心，最早开始时间取核心可用时间和所有前驱完成后同步、传输时间的最大值；同核与跨核等待代理分别为 100 与 1000 cycle，再加子图工作量代理得到预计完成时间。以上代理仅负责生成候选，不替代官方多核仿真；本文所有结果表的 Makespan 均来自官方评测。")
    add_equation(doc, r"EFT(S,c)=EST(S,c)+p(S)", "4")

    add_heading(doc, "4.3 完整依赖层窗口与分支保持", 2)
    add_paragraph(doc, "B2A 首先给算子标记依赖层：源算子为第 0 层，其余算子的层号比所有直接前驱的最大层号大 1。由于任一依赖边都指向更高层，按完整层切分不会把同层算子误认为前后依赖。实现对每层分别求 M、V 管线周期和，取两者较大值作为该层的切分权重，再按累计权重形成连续、非空的工作量均衡窗口。窗口数在 4、8、16 三个预设值中生成候选。")
    add_equation(doc, r"\ell(v)=1+\max_{u\in pred(v)}\ell(u)", "5")
    add_equation(doc, r"q_l=\max(\sum_{v\in V_l}c_M(v),\sum_{v\in V_l}c_V(v))", "6")
    add_paragraph(doc, "式（6）的逐层权重用于决定窗口边界；形成窗口后才另计算该窗口 M、V 管线周期总和的较大值作为窗口负载。两者一般不相等。每个窗口内部再忽略边的方向，只在窗口内的依赖边上寻找弱连通分支。分支内部已有依赖链，保留整条分支可避免任意切断同一局部分支；不同弱连通分支之间没有窗口内依赖边，因而可以按通信和负载进一步归组。")

    add_heading(doc, "4.4 边界张量感知打包与通信加权商图", 2)
    add_paragraph(doc, "对每个弱连通分支，统计其 M、V 管线工作量和跨出分支的前驱、后继张量 ID 集。窗口内若分支数量不超过 2k，则各分支独立成组；否则保留负载较大的 2k 个分支作为锚点，只把余下的小分支附到锚点组。选择附着位置时，先最大化共享前驱边界张量字节，再最大化共享后继边界张量字节，最后选择打包后双管线代理负载较小的组。这一字典序规则使局部张量接口与负载同时参与切图，且不拆开已识别出的弱连通分支。")
    add_paragraph(doc, "把每组视为一个子图后，构造商图：若原算子依赖跨组，则在对应子图之间连边。对同一商图边的张量按 ID 去重后求字节量，并加上直接算子依赖的携带量，得到 b(S,T)。式（3）使关键路径优先级包含预计的通信压力。排序后的就绪子图交由 4.2 节的共用 EFT 分核器处理，从而把 B2A 与 B0、B1 的比较集中在切图结构及由其形成的商图差异。程序还检查算子覆盖唯一性、窗口顺序和组间依赖，防止产生非法子图。")

    add_heading(doc, "4.5 官方评测择优与低核数方案继承", 2)
    add_paragraph(doc, "每个案例在 2—5 核分别生成 Single、B0、B1、B2A 的 4/8/16 窗口共六类候选。相同的正式方案经规范化后去重；有效方案均交由官方场景 A 评测。候选先比较 Makespan，再以附加搬运字节数打破并列，最后使用固定候选顺序保证选择可复现。单核方案单独作为 k=1 基准。")
    add_paragraph(doc, "由于增加核心数并不必然使重新切图后的方案更优，对 3—5 核还考察较低核数已经胜出的方案。继承时保留原算子到子图的映射及已有核心的任务顺序，只在目标核心数之后追加空核心序列；继承方案仍须重新进行合法性检查与官方评测。仅当其官方 Makespan 更小，或相同 Makespan 下附加搬运更少时，才替换当前方案。这样得到的改进应归因于候选继承与重新评测，不能混记为 B2A 切图单独带来的增益。")

    add_heading(doc, "4.6 实验结果与分析", 2)
    add_paragraph(doc, "对同一批 100 张输入图，1—5 核共有 500 个案例—核数组合，最终结果均由官方评测给出。输入图算子数从 766 到 38,666；原始 400 个多核组合均成功产出有效方案。表2给出最终方案的平均加速比、平均 Makespan 与平均附加搬运量，图1直观呈现加速比随核数变化的趋势。")
    result_table = []
    for k in range(1, 6):
        rows = by_core[k]
        result_table.append([
            str(k),
            f"{sum(float(r['final_speedup']) for r in rows)/100:.4f}",
            f"{sum(int(r['final_makespan']) for r in rows)/100:,.2f}",
            f"{sum(int(r['final_added_copy_bytes']) for r in rows)/100:,.2f}",
        ])
    add_table(doc, ["核心数", "平均加速比", "平均 Makespan/cycle", "平均附加搬运量/byte"], result_table, "表2 问题一最终官方评测汇总", font_size=12)
    add_figure(doc, "fig1_problem1_average_speedup.png", "图1 1—5 核平均加速比：最终方案与继承回退前的对照", 15.5)
    add_paragraph(doc, "从表2和图1可见，平均加速比在 2、3、4、5 核时分别为 1.6334、2.1050、2.5698、2.8905。平均值随核心数增加，但增量逐步受图依赖、跨子图数据周转及核心利用情况共同制约；图1本身只能显示结果趋势，不能单独量化各因素的因果贡献。图2进一步给出每个核心数下 100 个逐例加速比及其分布，表明同一核数下不同计算图的收益并不一致。")
    add_figure(doc, "fig2_problem1_case_speedup_distribution.png", "图2 2—5 核逐例加速比分布", 15.5)
    add_paragraph(doc, "在原始六类候选的 400 个多核组合中，B2A 三种窗口候选被选中 245 次；加入继承回退后的最终方案中，该计数为 220 次。被选中频次说明结构切图在此候选池中具有实际作用，但它并非单独的消融实验；所有候选共用 EFT 调度，且继承候选会改变终选方法的标签。因此本文不以该频次推断 B2A 的独立因果增益。")
    fallback_rows = []
    old_means = {int(r["cores"]): sum(float(x["speedup_t1_over_tk"]) for x in original_rows if int(x["cores"]) == int(r["cores"]))/100 for r in aggregate_rows if int(r["cores"]) > 1}
    for k in (3, 4, 5):
        a = aggregate[k]
        fallback_rows.append([
            str(k),
            str(sum(int(r["makespan_improvement"]) > 0 for r in by_core[k])),
            f"{old_means[k]:.4f}",
            f"{float(a['reported_arithmetic_mean_speedup']):.4f}",
            f"{float(a['reported_arithmetic_mean_speedup'])-old_means[k]:.4f}",
        ])
    add_table(doc, ["核心数", "严格改善组数", "回退前均值", "最终均值", "均值增量"], fallback_rows, "表3 继承回退的官方实测效果", font_size=12)
    add_paragraph(doc, "继承回退在 3、4、5 核分别严格缩短 3、9、24 个案例的 Makespan；对各核数全部 100 个案例求均值，平均加速比分别提高约 0.0049、0.0288、0.0814。36 个严格改善组合中，33 个组合的附加搬运量同时下降，其余 3 个持平。图3列出全部严格改善组合的时间降幅和搬运变化；例如 case_044 在 5 核时 Makespan 从 154,407 降为 87,589 cycle，附加搬运从 4,438,112 降为 2,915,104 byte。")
    add_figure(doc, "fig3_problem1_inherited_fallback_improvements.png", "图3 继承回退带来的 36 个严格改善组合", 12.2)
    add_paragraph(doc, "关于求解耗时，400 个原始多核记录中有 352 个标为本轮执行，其记录的总墙钟时间中位数为 13.572 秒；另外 48 个组采纳了先前开发阶段的记录。两类记录的运行来源不同，因此不合并为统一冷启动速度指标。上述有效性和性能结论均限于本题提供的 100 张计算图、固定配置及官方评测规则；B2A 窗口数和打包上限也来自本次固定候选设计，尚不能据此推断任意图规模下的最优参数。")

    doc.add_page_break()
    add_heading(doc, "附录 A 问题一逐例结果与复现说明", 1)
    add_paragraph(doc, "表 A1 与表 A2 分别列出同一 100 个案例在 1—5 核下的官方最终 Makespan 和总附加数据搬运量。两表与正文表2、图1—3使用同一批最终实测结果；加速比可由对应行的单核 Makespan 除以该核数 Makespan 复算。")
    makespan_data = [[r["case"], *[f"{int(r[f'makespan_k{k}']):,}" for k in range(1, 6)]] for r in appendix_rows]
    added_data = [[r["case"], *[f"{int(r[f'added_copy_bytes_k{k}']):,}" for k in range(1, 6)]] for r in appendix_rows]
    add_table(doc, ["案例", "1核", "2核", "3核", "4核", "5核"], makespan_data, "表 A1 100 个案例的 Makespan（cycle）", font_size=12)
    add_table(doc, ["案例", "1核", "2核", "3核", "4核", "5核"], added_data, "表 A2 100 个案例的附加搬运量（byte）", font_size=12)
    add_paragraph(doc, "问题一的求解入口为 run_problem1_full.py，低核数方案继承入口为 run_problem1_inherited_fallback.py；二者依赖 solver_problem1.py、candidate_manager_problem1.py 及 code 目录下的官方场景 A 校验与评测模块。输入位于 data/case_001.json 至 data/case_100.json，固定配置为 data/config.txt。结果整理脚本 artifacts/problem1_paper_figures/scripts/prepare_problem1_paper_data.py 从最终和回退前的五张结果表生成正文图数据与附录宽表。复现时应为新实验指定独立输出目录；继承回退的 artifact-root 参数指向实验基目录，full-root 参数可直接指定完整实验目录，候选计划与评测缓存目录按实际路径传入。")
    add_paragraph(doc, "参考文献：[1] TOPCUOGLU H, HARIRI S, WU M Y. Performance-effective and low-complexity task scheduling for heterogeneous computing[J]. IEEE Transactions on Parallel and Distributed Systems, 2002, 13(3): 260–274. DOI: 10.1109/71.993206.", indent=False, size=12)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "figures", 3, "appendix_cases", len(appendix_rows))


if __name__ == "__main__":
    main()
