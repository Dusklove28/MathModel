"""Read-only checks for the V2 review DOCX and frozen per-case evidence."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
DOCX = ROOT / "code/paper/draft/2026华为杯A题完整论文V2.docx"


def rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def wide_table(table, expected: dict[tuple[str, int], str], label: str) -> int:
    if len(table.rows) != 101 or len(table.columns) != 6:
        raise AssertionError(f"{label}: expected 100 cases and five core columns")
    checked = 0
    for row in table.rows[1:]:
        case = row.cells[0].text.strip()
        for cores in range(1, 6):
            observed = row.cells[cores].text.strip()
            desired = expected[(case, cores)]
            if observed != desired:
                raise AssertionError(f"{label}: {case}/{cores}: {observed!r} != {desired!r}")
            checked += 1
    return checked


def main() -> None:
    doc = Document(DOCX)
    q1 = rows(ROOT / "figures/q1/tables/appendix_problem1_wide.csv")
    q2 = rows(ROOT / "figures/q2/tables/selected_results.csv")
    q3 = rows(ROOT / "figures/q3/tables/case_core_2x2.csv")
    if (len(q1), len(q2), len(q3)) != (100, 500, 500):
        raise AssertionError("frozen source table cardinalities changed")
    if len(doc.tables) != 22 or len(doc.inline_shapes) != 10:
        raise AssertionError("V2 table or figure count changed unexpectedly")
    source_1 = {r["case"]: r for r in q1}
    source_2 = {(r["case"], int(r["cores"])): r for r in q2}
    source_3 = {(r["case"], int(r["cores"])): r for r in q3}
    checks = []
    for table_index, label, field in [
        (7, "A1", "makespan"), (8, "A2", "added_copy_bytes")
    ]:
        expected = {
            (case, k): f"{int(row[f'{field}_k{k}']):,}"
            for case, row in source_1.items() for k in range(1, 6)
        }
        checks.append(wide_table(doc.tables[table_index], expected, label))
    for table_index, label, field in [
        (9, "B1", "makespan_cycles"), (10, "B2", "added_copy_bytes"),
        (11, "B3", "spill_added_copy_bytes")
    ]:
        expected = {(case, k): f"{int(row[field]):,}"
                    for (case, k), row in source_2.items()}
        checks.append(wide_table(doc.tables[table_index], expected, label))
    q3_specs = [
        ("C1", "baseline_no_l2_makespan", "integer"),
        ("C2", "baseline_l2_makespan", "integer"),
        ("C3", "final_no_l2_makespan", "integer"),
        ("C4", "final_l2_makespan", "integer"),
        ("C5", "baseline_no_l2_added_copy_bytes", "integer"),
        ("C6", "baseline_l2_added_copy_bytes", "integer"),
        ("C7", "final_no_l2_added_copy_bytes", "integer"),
        ("C8", "final_l2_added_copy_bytes", "integer"),
        ("C9", "baseline_l2_cache_hit_rate_bytes", "percent"),
        ("C10", "final_l2_cache_hit_rate_bytes", "percent"),
    ]
    for offset, (label, field, kind) in enumerate(q3_specs):
        expected = {
            (case, k): (f"{int(row[field]):,}" if kind == "integer"
                        else f"{100*float(row[field]):.2f}")
            for (case, k), row in source_3.items()
        }
        checks.append(wide_table(doc.tables[12 + offset], expected, label))
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    figures = {int(m.group(1)) for p in paragraphs
               if (m := re.match(r"^图\s*(\d+)\s", p))}
    normal_tables = {int(m.group(1)) for p in paragraphs
                     if (m := re.match(r"^表\s*(\d+)\s", p))}
    letter_tables = {m.group(1) for p in paragraphs
                     if (m := re.match(r"^表\s*([ABC]\d+)\s", p))}
    if figures != set(range(1, 11)) or normal_tables != set(range(1, 7)):
        raise AssertionError("main figure/table captions are incomplete")
    if letter_tables != {"A1", "A2", "B1", "B2", "B3"} | {x[0] for x in q3_specs}:
        raise AssertionError("appendix table captions are incomplete")
    body = "\n".join(paragraphs[:paragraphs.index("参考文献")])
    forbidden = [x for x in (".py", "map_locality", "Gate A", "Stage 1", "artifact", "full-root")
                 if x.lower() in body.lower()]
    if forbidden:
        raise AssertionError(f"engineering jargon remains in main text: {forbidden}")
    print(json.dumps({
        "appendix_tables_checked": len(checks),
        "official_numeric_cells_checked": sum(checks),
        "problem3_numeric_cells_checked": sum(checks[5:]),
        "omml_equations": len(doc.element.xpath(".//m:oMath")),
        "figures": len(doc.inline_shapes),
        "review_only_ai_disclosure_pending": any("审阅待核" in p for p in paragraphs),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
