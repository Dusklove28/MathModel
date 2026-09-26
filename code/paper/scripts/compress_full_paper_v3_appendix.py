"""Keep every required per-case appendix metric while removing diagnostics.

V2 remains untouched.  The removed tables already exist in the tracked
Problem-2 and Problem-3 source CSV files; this script never recalculates data.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "code/paper/draft/2026华为杯A题完整论文V2.docx"
OUTPUT = ROOT / "code/paper/draft/2026华为杯A题完整论文V3-精简附录.docx"

DROP_CAPTIONS = (
    "表 B3 问题二最终 Spill（byte）",
    "表 C1 冻结方案无 L2 的 Makespan（cycle）",
    "表 C2 冻结方案有 L2 的 Makespan（cycle）",
    "表 C5 冻结方案无 L2 的总额外搬运（byte）",
    "表 C6 冻结方案有 L2 的总额外搬运（byte）",
    "表 C9 冻结方案有 L2 的字节命中率（%）",
)

REPLACE_PARAGRAPHS = {
    "下列表格按案例列出 1—5 核最终方案的官方 Makespan、新增搬运与 Spill。完整的候选失败记录、跨核传输和每核利用率保存在随附结果 CSV，表中不是只抽取改善案例。":
        "表 B1—B2 按案例列出 1—5 核最终方案的官方 Makespan 与总额外数据搬运量，"
        "覆盖全部 100 个案例。Spill、跨核传输、利用率和候选记录属于诊断性材料，"
        "保留在电子结果表 selected_results.csv，不与题目要求的两项逐例指标混排。",
    "附录 C 问题三 2×2 逐例官方结果":
        "附录 C 问题三最终方案的逐例官方结果",
    "C1—C4 列出冻结/最终方案在无/有 L2 条件下的逐例官方 Makespan；C5—C8 按相同四格列出总额外数据搬运量；C9—C10 分别列出两种有 L2 方案的按字节计 Cache 命中率。无 L2 条件不存在 Cache 命中率，记为不适用。每表覆盖同一 100 个案例、1—5 核，可按行复算两种配对比值；表中数值与正文使用同一份官方结果。":
        "表 C1—C2 列出最终方案在无 L2 和有 L2 条件下的逐例官方 Makespan；"
        "表 C3—C4 按相同两种配置列出总额外数据搬运量；表 C5 列出有 L2 条件下的"
        "按字节计 Cache 命中率。无 L2 条件没有 Cache 命中率，记为不适用。"
        "每表覆盖同一 100 个案例、1—5 核。同方案硬件效应可按行复算；"
        "用于 2×2 配对的冻结方案逐例数据完整保留在电子结果表 case_core_2x2.csv。",
    "表 C3 最终方案无 L2 的 Makespan（cycle）":
        "表 C1 最终方案无 L2 的 Makespan（cycle）",
    "表 C4 最终方案有 L2 的 Makespan（cycle）":
        "表 C2 最终方案有 L2 的 Makespan（cycle）",
    "表 C7 最终方案无 L2 的总额外搬运（byte）":
        "表 C3 最终方案无 L2 的总额外搬运（byte）",
    "表 C8 最终方案有 L2 的总额外搬运（byte）":
        "表 C4 最终方案有 L2 的总额外搬运（byte）",
    "表 C10 最终方案有 L2 的字节命中率（%）":
        "表 C5 最终方案有 L2 的字节命中率（%）",
    "逐例结果分别见 figures/q1/tables/appendix_problem1_wide.csv、figures/q2/tables/selected_results.csv 和 figures/q3/tables/case_core_2x2.csv。完整运行应从题目给定的 100 张输入图和统一配置开始；程序附件与论文 PDF 的命名、打包和上传由参赛队按当届系统要求最终确认。":
        "逐例原始结果分别见 figures/q1/tables/appendix_problem1_wide.csv、"
        "figures/q2/tables/selected_results.csv 和 figures/q3/tables/case_core_2x2.csv。"
        "附录 B 移出的逐例 Spill 见前一 CSV 的 spill_added_copy_bytes 列；"
        "附录 C 移出的冻结方案无/有 L2 的逐例时间、搬运与命中率，见后一 CSV 的 baseline_* 列。"
        "这两份文件是诊断和配对复核的电子材料；论文附录 A—C 仍列全题目要求的逐例指标。"
        "完整运行应从题目给定的 100 张输入图和统一配置开始；"
        "程序及电子数据附件的命名、打包和上传由参赛队按当届系统要求最终确认。",
}


def one_paragraph(doc: Document, text: str):
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text == text]
    if len(matches) != 1:
        raise AssertionError(f"expected exactly one paragraph {text!r}: {len(matches)}")
    return matches[0]


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    if OUTPUT.exists():
        raise FileExistsError(f"refusing to overwrite an existing review version: {OUTPUT}")
    doc = Document(SOURCE)
    if len(doc.tables) != 22 or len(doc.inline_shapes) != 10:
        raise AssertionError("V2 source table/figure count changed")

    for caption in DROP_CAPTIONS:
        paragraph = one_paragraph(doc, caption)
        caption_element = paragraph._element
        table_element = caption_element.getnext()
        if table_element is None or table_element.tag.rsplit("}", 1)[-1] != "tbl":
            raise AssertionError(f"table does not immediately follow {caption!r}")
        # Both nodes are removed as one unit, so no data table is left untitled.
        parent = caption_element.getparent()
        parent.remove(table_element)
        parent.remove(caption_element)

    for old, new in REPLACE_PARAGRAPHS.items():
        paragraph = one_paragraph(doc, old)
        if len(paragraph.runs) != 1:
            raise AssertionError(f"unexpected mixed run formatting: {old!r}")
        paragraph.runs[0].text = new

    if len(doc.tables) != 16 or len(doc.inline_shapes) != 10:
        raise AssertionError("V3 must have cover+body tables and nine appendix tables")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(json.dumps({
        "output": str(OUTPUT),
        "removed_diagnostic_tables": len(DROP_CAPTIONS),
        "retained_appendix_tables": 9,
        "all_figures_retained": len(doc.inline_shapes),
        "source_unchanged": str(SOURCE),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
